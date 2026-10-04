import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen slides (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # Blank slide

    # Theme Colors
    COLOR_BG = RGBColor(11, 15, 25)         # Dark navy/slate #0b0f19
    COLOR_CARD = RGBColor(18, 24, 38)       # #121826
    COLOR_PRIMARY = RGBColor(59, 130, 246)  # Electric Blue #3b82f6
    COLOR_ACCENT = RGBColor(16, 185, 129)   # Emerald Green #10b981
    COLOR_TITLE = RGBColor(248, 250, 252)   # #f8fafc
    COLOR_TEXT = RGBColor(148, 163, 184)    # #94a3b8
    COLOR_MUTED = RGBColor(100, 116, 139)   # #64748b
    COLOR_GOLD = RGBColor(245, 158, 11)     # Amber #f59e0b

    def add_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="BUS DEPOT MANAGEMENT SYSTEM"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    # Title & Subtitle box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(3.2))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "ACADEMIC CAPSTONE PROJECT"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Bus Depot Management System"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TITLE
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "A Role-Based Enterprise Web Platform for Fleet Logistics, Route Allocation & Duty Scheduling"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_TEXT

    # Meta Info Card (Placeholders for Student/College)
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.733), Inches(1.8))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_PRIMARY
    card.line.width = Pt(1)

    meta_box = s1.shapes.add_textbox(Inches(1.1), Inches(5.1), Inches(11.2), Inches(1.6))
    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True

    pm = tf_meta.paragraphs[0]
    pm.text = "PROJECT PRESENTATION DETAILS"
    pm.font.size = Pt(10)
    pm.font.bold = True
    pm.font.color.rgb = COLOR_GOLD
    pm.space_after = Pt(8)

    pm2 = tf_meta.add_paragraph()
    pm2.text = "• Student Name: [Your Name]      |   Roll Number: [Your Roll Number]      |   Department: Computer Science & Engineering\n• College / University: [Your College Name]      |   Project Guide: [Guide Name / Designation]"
    pm2.font.size = Pt(13)
    pm2.font.color.rgb = COLOR_TITLE

    # ==========================================
    # SLIDE 2: INTRODUCTION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s2)
    add_header(s2, "Introduction & Overview")

    def add_card(slide, x, y, w, h, title, items, icon="▶"):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = RGBColor(30, 41, 59)
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.2), Inches(w - 0.5), Inches(h - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        p.space_after = Pt(12)

        for item in items:
            pi = tf.add_paragraph()
            pi.text = f"•  {item}"
            pi.font.size = Pt(13)
            pi.font.color.rgb = COLOR_TEXT
            pi.space_after = Pt(6)

    add_card(s2, 0.8, 1.8, 5.6, 5.0, "The Modern Transit Challenge", [
        "Public and private bus transport networks operate hundreds of buses across complex city corridors.",
        "Daily depot activities require synchronized coordination between fleet maintenance, driver shifts, and timetable schedules.",
        "Manual registers and spreadsheets lead to scheduling conflicts, untracked delays, and delayed repairs.",
        "Depot supervisors need unified visibility to ensure optimal fleet turnaround and passenger safety."
    ], "🚌")

    add_card(s2, 6.8, 1.8, 5.6, 5.0, "Our Solution: Bus Depot Management System", [
        "A full-stack, enterprise-grade web application built using Python FastAPI and SQLite.",
        "Strict Role-Based Access Control (RBAC) across three distinct user roles: Admin, Operations, and Driver.",
        "Real-time asset tracking: Live availability counters for buses, drivers, routes, and dispatches.",
        "Full audit logging: Every login, vehicle assignment, schedule change, and status transition is recorded."
    ], "⚡")

    # ==========================================
    # SLIDE 3: PROBLEM STATEMENT
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s3)
    add_header(s3, "Problem Statement")

    add_card(s3, 0.8, 1.8, 3.6, 5.0, "Operational Friction", [
        "Manual pen-and-paper registers cause delays in morning vehicle dispatch.",
        "High risk of human error when assigning buses to routes.",
        "No live visibility into which buses are under maintenance vs road-ready."
    ], "⚠️")

    add_card(s3, 4.8, 1.8, 3.6, 5.0, "Lack of Authorization", [
        "Traditional desktop tools lack granular role permissions.",
        "Drivers and operations staff can accidentally modify core system settings.",
        "No audit trail to identify who made changes or unauthorized edits."
    ], "🔒")

    add_card(s3, 8.8, 1.8, 3.6, 5.0, "Communication Gaps", [
        "Drivers have no personal mobile-friendly portal to verify their daily roster.",
        "Operations managers cannot track trip completion in real-time.",
        "Depot administrators struggle to generate fleet performance analytics."
    ], "📡")

    # ==========================================
    # SLIDE 4: EXISTING VS PROPOSED SYSTEM
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s4)
    add_header(s4, "Existing System vs. Proposed System")

    add_card(s4, 0.8, 1.8, 5.6, 5.0, "Existing Legacy System", [
        "Physical ledger registers or static excel spreadsheets.",
        "Double-booking of buses or drivers due to desynchronized data.",
        "Zero password security: plain records stored without encryption.",
        "No automated release of vehicles once trips reach completion.",
        "Manual calculation of reports and maintenance tracking.",
        "High turnaround time for incident resolution."
    ], "❌")

    add_card(s4, 6.8, 1.8, 5.6, 5.0, "Proposed Web Management System", [
        "Centralized SQLite relational database with ACID compliance.",
        "FastAPI asynchronous backend with sub-millisecond API response.",
        "Bcrypt password hashing + JWT session security + Cookie protection.",
        "Automated lifecycle: completing a trip instantly frees bus & driver.",
        "Visual analytics: Fleet breakdown, driver duty status, route loads.",
        "Dedicated self-service portal for drivers to update live trip states."
    ], "✅")

    # ==========================================
    # SLIDE 5: PROJECT OBJECTIVES
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s5)
    add_header(s5, "Project Objectives")

    add_card(s5, 0.8, 1.8, 5.6, 2.3, "1. Role-Based Access Control", [
        "Enforce strict server-side RBAC separating Admin, Operations, and Driver portals.",
        "Guarantee URL tampering redirects unprivileged users immediately."
    ], "🎯")

    add_card(s5, 6.8, 1.8, 5.6, 2.3, "2. End-to-End Fleet Logistics", [
        "Manage complete bus lifecycle from registration to periodic maintenance.",
        "Track capacity, bus type (AC, Express, Electric), and roadworthiness."
    ], "🚌")

    add_card(s5, 0.8, 4.4, 5.6, 2.3, "3. Intelligent Dispatching", [
        "Prevent assigning unavailable or maintenance-bound buses to new trips.",
        "Automatically transition driver rosters between Active, On Duty, and Off Duty."
    ], "⚙️")

    add_card(s5, 6.8, 4.4, 5.6, 2.3, "4. Real-time Audit & Analytics", [
        "Capture every security and operational event in an immutable activity log.",
        "Provide interactive visual charts for fleet utilization and route loads."
    ], "📊")

    # ==========================================
    # SLIDE 6: TECHNOLOGY STACK
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s6)
    add_header(s6, "Technology Stack & Rationale")

    add_card(s6, 0.8, 1.8, 3.6, 5.0, "Backend (Python & FastAPI)", [
        "Python 3.12: High productivity, robust typing, and extensive library ecosystem.",
        "FastAPI: Modern, async-native ASGI web framework with automatic OpenAPI docs.",
        "SQLAlchemy 2.0 ORM: Safe object-relational mapping preventing SQL injection.",
        "Uvicorn: Lightning-fast ASGI production server."
    ], "🐍")

    add_card(s6, 4.8, 1.8, 3.6, 5.0, "Security & Database", [
        "SQLite 3: Zero-configuration, zero-latency serverless relational storage.",
        "Bcrypt: Salting & adaptive hashing algorithm against dictionary attacks.",
        "PyJWT: Stateless signed authentication tokens with expiration timestamps.",
        "HTTP-Only Cookies: Mitigation against Cross-Site Scripting (XSS) attacks."
    ], "🛡️")

    add_card(s6, 8.8, 1.8, 3.6, 5.0, "Frontend (HTML5, CSS3, JS)", [
        "HTML5 Semantic Architecture: Accessible tables, forms, and dialog wrappers.",
        "Vanilla CSS3: Custom enterprise design system with dark mode & glassmorphism.",
        "Vanilla ES6+ JavaScript: Fetch API async/await controllers with real-time DOM updates.",
        "Zero Heavy Framework Overhead: Instant load time and zero build step required."
    ], "💻")

    # ==========================================
    # SLIDE 7: SYSTEM ARCHITECTURE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s7)
    add_header(s7, "Multi-Tier System Architecture")

    # Architecture diagram cards
    add_card(s7, 0.8, 1.8, 2.6, 5.0, "Presentation Tier", [
        "Responsive Web UI",
        "Admin Dashboard",
        "Operations Portal",
        "Driver Mobile View",
        "Vanilla JS Controllers",
        "Fetch API Async Client"
    ], "1️⃣")

    add_card(s7, 3.7, 1.8, 2.8, 5.0, "Security & Gateway", [
        "HTTP-only Cookies",
        "Bearer JWT Tokens",
        "Password Verification",
        "Role Guard Middleware",
        "CORS & Input Validation",
        "Session Expiry Checks"
    ], "2️⃣")

    add_card(s7, 6.8, 1.8, 2.8, 5.0, "Application Tier", [
        "FastAPI Async Routers",
        "Auth Service (/auth)",
        "Admin Engine (/admin)",
        "Ops Engine (/operations)",
        "Driver Engine (/driver)",
        "Activity Audit Logger"
    ], "3️⃣")

    add_card(s7, 9.9, 1.8, 2.6, 5.0, "Data Tier", [
        "SQLAlchemy 2.0 ORM",
        "SQLite Relational DB",
        "Users & Drivers",
        "Buses & Routes",
        "Trips & Assignments",
        "Audit Activity Logs"
    ], "4️⃣")

    # ==========================================
    # SLIDE 8: MODULES & USER ROLES
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s8)
    add_header(s8, "Modules & Role-Based Privileges (RBAC)")

    add_card(s8, 0.8, 1.8, 3.6, 5.0, "1. Administrator", [
        "Total Depot Oversight",
        "User Management (Create, Edit, Deactivate, Roles)",
        "Bus Master Catalog & Maintenance Logs",
        "Driver Registration & License Management",
        "Transit Routes & Stop Configuration",
        "Depot Analytics & Visual Fleet Reports",
        "Full Security Audit Activity Trail"
    ], "👑")

    add_card(s8, 4.8, 1.8, 3.6, 5.0, "2. Operations Staff", [
        "Fleet Dispatch Operations",
        "View & Filter Real-time Bus Status",
        "Check Driver Availability & Roster",
        "Create Trips & Assign Bus + Driver",
        "Manage Operational Assignments",
        "Override Trip Status When Permitted",
        "Restricted: Cannot alter users/roles"
    ], "📋")

    add_card(s8, 8.8, 1.8, 3.6, 5.0, "3. Driver Portal", [
        "Personalized Pilot View",
        "View Assigned Bus Specs & Status",
        "View Active Route & Corridor Stops",
        "View Personal Duty Schedule",
        "Start Trip (Assigned ➔ In Progress)",
        "Complete Trip (Auto-releases vehicle)",
        "Restricted: No administrative access"
    ], "👨‍✈️")

    # ==========================================
    # SLIDE 9: AUTHENTICATION & SECURITY FLOW
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s9)
    add_header(s9, "Authentication & Security Architecture")

    add_card(s9, 0.8, 1.8, 5.6, 5.0, "Login & Token Verification Flow", [
        "Step 1: User submits Username/Email and Password via HTTPS.",
        "Step 2: Backend fetches user record by unique identifier.",
        "Step 3: Bcrypt verifies submitted password against stored hash.",
        "Step 4: If valid, backend generates signed HS256 JWT access token.",
        "Step 5: Token set in HTTP-Only SameSite cookie and returned in JSON.",
        "Step 6: Backend dynamically routes user to role-specific dashboard.",
        "Step 7: Every sub-request validates signature and active role."
    ], "🔐")

    add_card(s9, 6.8, 1.8, 5.6, 5.0, "Security Safeguards Implemented", [
        "✓ Zero Plain-Text Passwords: All credentials hashed with 12 rounds of bcrypt.",
        "✓ Server-Side RBAC: Endpoints protected by FastAPI Depends(require_role).",
        "✓ URL Tampering Proof: Visiting unauthorized dashboard URL redirects with warning.",
        "✓ SQL Injection Prevention: Parameterized SQLAlchemy queries.",
        "✓ XSS Protection: Strict template escaping & HTTP-only cookies.",
        "✓ Audit Trail: Every failed login, creation, and dispatch logged."
    ], "🛡️")

    # ==========================================
    # SLIDE 10: DATABASE SCHEMA & ENTITIES
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s10)
    add_header(s10, "Relational Database Design (SQLite)")

    add_card(s10, 0.8, 1.8, 3.6, 5.0, "Identity & Personnel", [
        "users: id, username, email, password_hash, role, is_active, created_at.",
        "drivers: id, user_id (FK), driver_id, name, phone, license_number, status.",
        "One-to-One link between driver user account and personnel record."
    ], "👥")

    add_card(s10, 4.8, 1.8, 3.6, 5.0, "Fleet & Transit", [
        "buses: id, bus_number, registration_number, bus_type, capacity, status, maintenance_status.",
        "routes: id, route_number, source, destination, stops, status.",
        "Enforces unique bus numbers and registration plates."
    ], "🚌")

    add_card(s10, 8.8, 1.8, 3.6, 5.0, "Operations & Audit", [
        "trips: id, route_id (FK), bus_id (FK), driver_id (FK), departure_time, arrival_time, status.",
        "assignments: id, trip_id (FK), bus_id, driver_id, assigned_by, notes, status.",
        "activity_logs: id, user_id (FK), username, role, action, details, timestamp."
    ], "📜")

    # ==========================================
    # SLIDE 11: KEY APPLICATION FEATURES & UI
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s11)
    add_header(s11, "Application Highlights & Key Features")

    add_card(s11, 0.8, 1.8, 5.6, 5.0, "Enterprise UI / UX Features", [
        "Responsive Dashboard Shell: Collapsible sidebar, topbar with live clock and role indicator.",
        "Single-Page Transitions: Seamless switching between views without page refreshes.",
        "Real-Time Statistics: Metrics cards with visual counters and color-coded status badges.",
        "Search & Filtering: Instant client-side search across buses, drivers, routes, and logs.",
        "Interactive Modal Windows: Clean popup forms for creating and editing records.",
        "Animated Toast Alerts: Floating non-intrusive notification feedback."
    ], "✨")

    add_card(s11, 6.8, 1.8, 5.6, 5.0, "Automated Operational Workflow", [
        "Intelligent Assignment: Automatically marks bus 'Assigned' and driver 'On Duty'.",
        "Vehicle Protection: Prevents assigning buses marked 'Under Maintenance'.",
        "Driver Self-Service: Drivers click 'Start Trip' to mark route in-progress.",
        "Automated Resource Release: Completing a trip frees the bus back to 'Available'!",
        "Comprehensive Audit Log: Retains timestamped record of every administrative action.",
        "100% Automated Test Suite: 29 out of 29 end-to-end unit tests passed."
    ], "🚀")

    # ==========================================
    # SLIDE 12: CONCLUSION & FUTURE SCOPE
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s12)
    add_header(s12, "Conclusion & Future Enhancements")

    add_card(s12, 0.8, 1.8, 5.6, 5.0, "Project Conclusion", [
        "Successfully engineered a complete, production-ready Bus Depot Management System.",
        "Implemented robust, secure authentication and role-based authorization.",
        "Automated operational workflows, reducing vehicle dispatch bottlenecks.",
        "Eliminated paperwork and spreadsheet dependencies with centralized SQLite storage.",
        "Designed a professional, responsive user interface suitable for depot scale."
    ], "🎓")

    add_card(s12, 6.8, 1.8, 5.6, 5.0, "Future Scope & Enhancements", [
        "GPS Fleet Tracking: Real-time telemetry via IoT OBD-II dongles on buses.",
        "Automated Passenger Ticketing: Integration with smart cards / QR code tickets.",
        "Predictive Maintenance: AI algorithms to forecast mechanical failures from mileage.",
        "Automated Crew Rostering: Optimization algorithms for driver duty scheduling.",
        "Passenger Mobile App: Real-time bus arrival times and seat availability."
    ], "🔮")

    # Output Presentation
    output_path = "/home/rkvalley-rgukt/Desktop/login/Bus_Depot_Management_System.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
